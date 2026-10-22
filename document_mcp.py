
from pathlib import Path

import fitz  # PyMuPDF
from openpyxl import load_workbook
from mcp.server.fastmcp import FastMCP


# --------------------------------------------------
# Configuration
# --------------------------------------------------

DOCUMENT_DIR = Path(r"D:\OpenCodeDocuments").resolve()

# Maximum extracted PDF text returned by one tool call.
# If a PDF is larger, use read_pdf with page ranges.
MAX_TEXT_CHARS = 40000

mcp = FastMCP("local-documents")

PDF_EXTENSIONS = {".pdf"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}


# --------------------------------------------------
# Shared helpers
# --------------------------------------------------

def resolve_file(file_name: str, allowed_extensions: set[str]) -> Path:
    """Find a permitted file inside the document directory."""

    DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)

    if not file_name or Path(file_name).name != file_name:
        raise ValueError(
            "Provide only the file name, not a path. "
            "Example: incident-report.pdf"
        )

    path = (DOCUMENT_DIR / file_name).resolve()

    if not path.is_relative_to(DOCUMENT_DIR):
        raise ValueError("Access outside the document directory is denied.")

    if path.suffix.lower() not in allowed_extensions:
        raise ValueError(
            f"Unsupported file type. Allowed extensions: "
            f"{', '.join(sorted(allowed_extensions))}"
        )

    if not path.is_file():
        raise FileNotFoundError(
            f"File '{file_name}' was not found in {DOCUMENT_DIR}"
        )

    return path


# --------------------------------------------------
# Tool 1: List available documents
# --------------------------------------------------

@mcp.tool()
def list_documents() -> dict:
    """List PDF and Excel documents available to OpenCode."""

    DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)

    files = sorted(
        file.name
        for file in DOCUMENT_DIR.iterdir()
        if file.is_file()
        and file.suffix.lower() in PDF_EXTENSIONS | EXCEL_EXTENSIONS
    )

    return {
        "directory": str(DOCUMENT_DIR),
        "count": len(files),
        "files": files,
    }


# --------------------------------------------------
# Tool 2: Read PDF pages
# --------------------------------------------------

@mcp.tool()
def read_pdf(
    file_name: str,
    start_page: int = 1,
    end_page: int = 0,
) -> dict:
    """
    Extract text from a PDF.

    Reads all pages by default. Use end_page=0 for the last page.
    For large documents, specify a smaller page range.
    Page numbers start at 1.
    """

    if start_page < 1 or end_page < 0:
        raise ValueError("Page numbers must be positive; end_page can be 0.")

    path = resolve_file(file_name, PDF_EXTENSIONS)

    sections = []
    remaining = MAX_TEXT_CHARS
    last_page_read = start_page - 1
    truncated = False

    with fitz.open(path) as pdf:
        total_pages = len(pdf)

        if total_pages == 0:
            return {
                "file": path.name,
                "total_pages": 0,
                "text": "",
                "note": "The PDF contains no pages.",
            }

        if start_page > total_pages:
            raise ValueError(
                f"This PDF has {total_pages} pages. "
                f"Choose a start_page between 1 and {total_pages}."
            )

        actual_end = (
            total_pages
            if end_page == 0
            else min(end_page, total_pages)
        )

        if actual_end < start_page:
            raise ValueError("end_page must be greater than or equal to start_page.")

        for page_number in range(start_page, actual_end + 1):
            page_text = pdf[page_number - 1].get_text("text")

            if len(page_text) > remaining:
                page_text = page_text[:remaining]
                truncated = True

            sections.append(
                f"--- Page {page_number} ---\n{page_text}"
            )
            last_page_read = page_number
            remaining -= len(page_text)

            if truncated or remaining <= 0:
                if page_number < actual_end or truncated:
                    truncated = True
                break

    return {
        "file": path.name,
        "total_pages": total_pages,
        "requested_pages": f"{start_page}-{actual_end}",
        "pages_read": f"{start_page}-{last_page_read}",
        "text": "\n\n".join(sections),
        "truncated": truncated,
        "note": (
            "Text was truncated at the output limit. "
            "Read smaller page ranges to retrieve the remaining content."
            if truncated
            else
            "If a page has no extracted text, it may be scanned and require OCR."
        ),
    }


# --------------------------------------------------
# Tool 3: Search PDF text
# --------------------------------------------------

@mcp.tool()
def search_pdf(
    file_name: str,
    query: str,
    max_results: int = 20,
) -> dict:
    """
    Search extracted text in a PDF and return matching pages
    with short snippets. This searches text, not scanned images.
    """

    if not query or not query.strip():
        raise ValueError("Provide a search term.")

    if not 1 <= max_results <= 100:
        raise ValueError("max_results must be between 1 and 100.")

    path = resolve_file(file_name, PDF_EXTENSIONS)
    matches = []
    needle = query.casefold()

    with fitz.open(path) as pdf:
        total_pages = len(pdf)

        for page_index, page in enumerate(pdf):
            page_text = page.get_text("text")
            lower_text = page_text.casefold()
            search_from = 0

            while len(matches) < max_results:
                position = lower_text.find(needle, search_from)

                if position == -1:
                    break

                snippet_start = max(0, position - 150)
                snippet_end = min(
                    len(page_text),
                    position + len(query) + 250,
                )

                matches.append({
                    "page": page_index + 1,
                    "snippet": page_text[snippet_start:snippet_end].strip(),
                })

                search_from = position + len(needle)

            if len(matches) >= max_results:
                break

    return {
        "file": path.name,
        "total_pages": total_pages,
        "query": query,
        "match_count": len(matches),
        "matches": matches,
        "note": "No matches may mean the term is absent or the page needs OCR.",
    }


# --------------------------------------------------
# Tool 4: Inspect Excel workbook
# --------------------------------------------------

@mcp.tool()
def inspect_excel(
    file_name: str,
    max_rows: int = 10,
    max_columns: int = 30,
) -> dict:
    """
    Inspect workbook sheet names, dimensions, headers and sample rows.
    Formula cells are returned as formulas, not calculated values.
    """

    if not 1 <= max_rows <= 100:
        raise ValueError("max_rows must be between 1 and 100.")

    if not 1 <= max_columns <= 100:
        raise ValueError("max_columns must be between 1 and 100.")

    path = resolve_file(file_name, EXCEL_EXTENSIONS)
    workbook = load_workbook(path, read_only=True, data_only=False)

    sheets = []

    try:
        for worksheet in workbook.worksheets:
            sample = []

            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(max_rows, worksheet.max_row or 1),
                max_col=min(max_columns, worksheet.max_column or 1),
                values_only=True,
            ):
                sample.append(list(row))

            sheets.append({
                "sheet_name": worksheet.title,
                "rows": worksheet.max_row,
                "columns": worksheet.max_column,
                "sample_rows": sample,
            })
    finally:
        workbook.close()

    return {
        "file": path.name,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "note": (
            "Formula cells are shown as formulas. "
            "Openpyxl does not calculate formulas."
        ),
    }


# --------------------------------------------------
# Tool 5: Read Excel rows
# --------------------------------------------------

@mcp.tool()
def read_excel_rows(
    file_name: str,
    sheet_name: str,
    start_row: int = 1,
    end_row: int = 20,
    max_columns: int = 50,
) -> dict:
    """
    Read a selected row range from an Excel worksheet.
    Row numbers start at 1. Maximum 100 rows per call.
    """

    if start_row < 1 or end_row < start_row:
        raise ValueError("Provide a valid row range starting at row 1 or later.")

    if end_row - start_row + 1 > 100:
        raise ValueError("Read at most 100 rows per call.")

    if not 1 <= max_columns <= 100:
        raise ValueError("max_columns must be between 1 and 100.")

    path = resolve_file(file_name, EXCEL_EXTENSIONS)
    workbook = load_workbook(path, read_only=True, data_only=False)

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. "
                f"Available sheets: {', '.join(workbook.sheetnames)}"
            )

        worksheet = workbook[sheet_name]
        last_row = min(end_row, worksheet.max_row or 0)
        last_column = min(max_columns, worksheet.max_column or 0)

        rows = []

        if last_row >= start_row and last_column > 0:
            for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=start_row,
                    max_row=last_row,
                    min_col=1,
                    max_col=last_column,
                    values_only=True,
                ),
                start=start_row,
            ):
                rows.append({
                    "row_number": row_number,
                    "values": list(row),
                })
    finally:
        workbook.close()

    return {
        "file": path.name,
        "sheet_name": sheet_name,
        "requested_rows": f"{start_row}-{end_row}",
        "rows_returned": len(rows),
        "rows": rows,
        "note": "Formula cells are returned as formulas, not calculated values.",
    }


# --------------------------------------------------
# Start MCP server
# --------------------------------------------------

if __name__ == "__main__":
    mcp.run(transport="stdio")
